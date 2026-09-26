"""Type sniffing and text extraction for uploaded or fetched documents.

The declared content type is only a hint: the file's bytes decide the type.
PDF pages with little or no text layer (scanned pages, figure-only slides) are
flagged for vision OCR and can be rendered to JPEG with pypdfium2.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

MAX_PDF_PAGES = 400
LOW_TEXT_CHARS = 80
FIGURE_PAGE_TEXT_CHARS = 400
RENDER_MAX_SIDE_PX = 1600
MAX_IMAGE_PIXELS = 50_000_000
MAX_ZIP_ENTRIES = 5000
MAX_ZIP_UNCOMPRESSED = 200 * 1024 * 1024

CONTENT_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "markdown": "text/markdown",
    "text": "text/plain",
    "html": "text/html",
    "png": "image/png",
    "jpeg": "image/jpeg",
    "webp": "image/webp",
}
IMAGE_KINDS = {"png", "jpeg", "webp"}


class ExtractionError(ValueError):
    """The file is unsupported, corrupt or unreadable. Message is safe to show."""


@dataclass(frozen=True)
class PageText:
    number: int | None
    text: str
    needs_vision: bool = False


@dataclass(frozen=True)
class ExtractedDocument:
    kind: str
    content_type: str
    title: str | None
    pages: list[PageText]


def sniff_kind(data: bytes, filename: str, declared_type: str | None) -> str:
    """Decide the document kind from magic bytes; the name/header only disambiguate text."""
    head = data[:2048]
    if b"%PDF-" in head[:1024]:
        return "pdf"
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if head.startswith(b"\xff\xd8\xff"):
        return "jpeg"
    if head[:4] == b"RIFF" and head[8:12] == b"WEBP":
        return "webp"
    if head.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if "word/document.xml" in {entry.filename for entry in entries}:
                    # Zip-bomb guard: python-docx inflates parts in memory.
                    if len(entries) > MAX_ZIP_ENTRIES or sum(entry.file_size for entry in entries) > MAX_ZIP_UNCOMPRESSED:
                        raise ExtractionError("the Word document is too large once decompressed")
                    return "docx"
        except zipfile.BadZipFile as exc:
            raise ExtractionError("the archive is corrupt") from exc
        raise ExtractionError("only Word .docx archives are supported")
    if b"\x00" in data[:8192] or _decode(data) is None:
        raise ExtractionError("unsupported file type; upload PDF, DOCX, Markdown, text, HTML, PNG, JPEG or WebP")
    suffix = Path(filename).suffix.lower()
    declared = (declared_type or "").split(";")[0].strip().lower()
    lowered = data[:1024].decode("utf-8", errors="ignore").lower()
    if suffix in {".html", ".htm"} or declared in {"text/html", "application/xhtml+xml"} \
            or "<!doctype html" in lowered or "<html" in lowered:
        return "html"
    if suffix in {".md", ".markdown"} or declared == "text/markdown":
        return "markdown"
    return "text"


def extract(data: bytes, kind: str) -> ExtractedDocument:
    if kind == "pdf":
        return _pdf(data)
    if kind == "docx":
        return _docx(data)
    if kind == "html":
        return _html(data)
    if kind in {"markdown", "text"}:
        text = _normalize(_decode(data) or "")
        return ExtractedDocument(kind, CONTENT_TYPES[kind], None, [PageText(None, text)])
    if kind in IMAGE_KINDS:
        validate_image(data)
        return ExtractedDocument(kind, CONTENT_TYPES[kind], None, [PageText(1, "", needs_vision=True)])
    raise ExtractionError("unsupported file type")


def _decode(data: bytes) -> str | None:
    for encoding in ("utf-8-sig", "utf-16") if data[:2] in {b"\xff\xfe", b"\xfe\xff"} else ("utf-8-sig",):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    try:
        text = data.decode("cp1252")
    except UnicodeDecodeError:
        return None
    control = sum(1 for char in text[:4000] if ord(char) < 32 and char not in "\n\r\t\f")
    return text if control < 10 else None


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\f", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _pdf(data: bytes) -> ExtractedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractionError("password-protected PDFs are not supported")
        page_list = list(reader.pages)
    except ExtractionError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError, OSError) as exc:
        raise ExtractionError("could not read this PDF") from exc
    if not page_list:
        raise ExtractionError("the PDF has no pages")
    if len(page_list) > MAX_PDF_PAGES:
        raise ExtractionError(f"PDFs are limited to {MAX_PDF_PAGES} pages")
    pages: list[PageText] = []
    for number, page in enumerate(page_list, start=1):
        try:
            text = _normalize(page.extract_text() or "")
        except Exception:  # pypdf raises many types on malformed content streams
            text = ""
        images = _image_count(page)
        compact = len(re.sub(r"\s+", "", text))
        needs_vision = compact < LOW_TEXT_CHARS or (images > 0 and compact < FIGURE_PAGE_TEXT_CHARS)
        pages.append(PageText(number, text, needs_vision))
    title = None
    try:
        metadata = reader.metadata
        title = str(metadata.title).strip()[:300] if metadata and metadata.title else None
    except Exception:
        title = None
    return ExtractedDocument("pdf", CONTENT_TYPES["pdf"], title, pages)


def _image_count(page) -> int:
    try:
        resources = page.get("/Resources")
        resources = resources.get_object() if resources is not None else None
        objects = resources.get("/XObject") if resources else None
        objects = objects.get_object() if objects is not None else {}
        return sum(1 for value in objects.values() if value.get_object().get("/Subtype") == "/Image")
    except Exception:
        return 0


def render_pdf_page(data: bytes, page_number: int) -> bytes:
    """Render one 1-based PDF page to a bounded-size JPEG for vision OCR."""
    import pypdfium2 as pdfium

    try:
        document = pdfium.PdfDocument(data)
        try:
            page = document[page_number - 1]
            width, height = page.get_size()
            scale = max(0.5, min(2.5, RENDER_MAX_SIDE_PX / max(width, height, 1)))
            image = page.render(scale=scale).to_pil()
        finally:
            document.close()
    except (pdfium.PdfiumError, IndexError, ValueError) as exc:
        raise ExtractionError(f"could not render PDF page {page_number}") from exc
    return _jpeg(image)


def validate_image(data: bytes) -> tuple[int, int]:
    from PIL import Image, UnidentifiedImageError

    Image.MAX_IMAGE_PIXELS = MAX_IMAGE_PIXELS
    try:
        with Image.open(io.BytesIO(data)) as image:
            image.verify()
            return image.size
    except (UnidentifiedImageError, Image.DecompressionBombError, OSError, SyntaxError) as exc:
        raise ExtractionError("the image is corrupt or too large") from exc


def prepare_image(data: bytes) -> bytes:
    """Downscale an uploaded image to a bounded JPEG (cheaper vision calls)."""
    from PIL import Image, UnidentifiedImageError

    Image.MAX_IMAGE_PIXELS = MAX_IMAGE_PIXELS
    try:
        with Image.open(io.BytesIO(data)) as image:
            image.load()
            return _jpeg(image)
    except (UnidentifiedImageError, Image.DecompressionBombError, OSError) as exc:
        raise ExtractionError("the image is corrupt or too large") from exc


def _jpeg(image) -> bytes:
    from PIL import Image

    if image.mode not in {"RGB", "L"}:
        background = Image.new("RGB", image.size, "white")
        converted = image.convert("RGBA")
        background.paste(converted, mask=converted.split()[-1])
        image = background
    image.thumbnail((RENDER_MAX_SIDE_PX, RENDER_MAX_SIDE_PX))
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", quality=85, optimize=True)
    return buffer.getvalue()


def _docx(data: bytes) -> ExtractedDocument:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises KeyError/ValueError/lxml errors
        raise ExtractionError("could not read this Word document") from exc
    lines: list[str] = []
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            lines.append(_docx_paragraph(Paragraph(element, document)))
        elif tag == "tbl":
            table = Table(element, document)
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            lines.append("\n".join(row for row in rows if row.strip(" |")))
    title = None
    try:
        title = (document.core_properties.title or "").strip()[:300] or None
    except Exception:
        title = None
    text = _normalize("\n\n".join(line for line in lines if line.strip()))
    return ExtractedDocument("docx", CONTENT_TYPES["docx"], title, [PageText(None, text)])


def _docx_paragraph(paragraph) -> str:
    text = paragraph.text.strip()
    if not text:
        return ""
    style = (paragraph.style.name if paragraph.style is not None else "") or ""
    match = re.match(r"Heading (\d)", style)
    if match:
        return f"{'#' * min(6, int(match.group(1)))} {text}"
    if style == "Title":
        return f"# {text}"
    if style.startswith("List"):
        return f"- {text}"
    return text


class _HtmlText(HTMLParser):
    _SKIP = {"script", "style", "noscript", "svg", "nav", "footer", "form", "iframe", "template", "head"}
    _BLOCK = {"p", "div", "section", "article", "main", "br", "tr", "table", "ul", "ol", "blockquote", "pre", "header"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.skip_depth = 0
        self.title: str | None = None
        self._in_title = False

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag == "title":
            self._in_title = True
        if tag in self._SKIP:
            self.skip_depth += 1
        elif self.skip_depth == 0 and re.fullmatch(r"h[1-6]", tag):
            self.parts.append(f"\n\n{'#' * int(tag[1])} ")
        elif self.skip_depth == 0 and tag == "li":
            self.parts.append("\n- ")
        elif self.skip_depth == 0 and tag in self._BLOCK:
            self.parts.append("\n\n")

    def handle_endtag(self, tag: str) -> None:
        if tag == "title":
            self._in_title = False
        if tag in self._SKIP and self.skip_depth:
            self.skip_depth -= 1
        elif self.skip_depth == 0 and (re.fullmatch(r"h[1-6]", tag) or tag in self._BLOCK):
            self.parts.append("\n\n")

    def handle_data(self, data: str) -> None:
        if self._in_title and not self.title:
            self.title = " ".join(data.split())[:300] or None
        if self.skip_depth == 0 and data:
            collapsed = " ".join(data.split())
            if not collapsed:
                self.parts.append(" ")
                return
            # Keep word boundaries around inline tags: "<i>fast</i> ones" -> "fast ones".
            self.parts.append((" " if data[0].isspace() else "") + collapsed + (" " if data[-1].isspace() else ""))


def _html(data: bytes) -> ExtractedDocument:
    parser = _HtmlText()
    try:
        parser.feed(_decode(data) or "")
        parser.close()
    except (AssertionError, ValueError) as exc:
        raise ExtractionError("could not read this HTML page") from exc
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in "".join(parser.parts).split("\n")]
    text = _normalize("\n".join(lines))
    return ExtractedDocument("html", CONTENT_TYPES["html"], parser.title, [PageText(None, text)])
