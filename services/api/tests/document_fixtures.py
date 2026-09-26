"""Tiny generated document fixtures and a fake OpenAI-compatible provider for document tests."""

from __future__ import annotations

import io
import json
from datetime import UTC, datetime
from uuid import uuid4

from meetings_contracts import EmbeddingResult, TextGenerationResult


def text_pdf(pages: list[str]) -> bytes:
    """A valid PDF whose pages carry a real text layer (empty string = page without text)."""
    objects: list[bytes] = []
    kids = []
    font_id = 3 + 2 * len(pages)
    for index, body in enumerate(pages):
        page_id, content_id = 3 + 2 * index, 4 + 2 * index
        kids.append(f"{page_id} 0 R")
        lines = "".join(f"({line}) Tj 0 -16 Td " for line in body.split("\n")) if body else ""
        stream = f"BT /F1 12 Tf 72 720 Td {lines}ET".encode() if body else b""
        objects.append(f"{page_id} 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                       f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >> endobj".encode())
        objects.append(f"{content_id} 0 obj << /Length {len(stream)} >> stream\n".encode() + stream + b"\nendstream endobj")
    header = [b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj",
              f"2 0 obj << /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >> endobj".encode()]
    objects = header + objects + [f"{font_id} 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj".encode()]
    output = bytearray(b"%PDF-1.4\n")
    offsets = []
    for item in objects:
        offsets.append(len(output))
        output += item + b"\n"
    xref = len(output)
    output += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    output += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    output += f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(output)


def image_bytes(kind: str = "PNG", size: tuple[int, int] = (320, 200)) -> bytes:
    from PIL import Image, ImageDraw

    image = Image.new("RGB", size, "white")
    ImageDraw.Draw(image).text((10, 10), "INVOICE 42", fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format=kind)
    return buffer.getvalue()


def scanned_pdf() -> bytes:
    """An image-only PDF (what a scanner produces): no text layer at all."""
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (600, 800), "white")
    ImageDraw.Draw(image).text((40, 40), "Scanned contract", fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PDF")
    return buffer.getvalue()


def docx_bytes() -> bytes:
    from docx import Document

    document = Document()
    document.core_properties.title = "Service catalogue"
    document.add_heading("Services", level=1)
    document.add_paragraph("We build retrieval platforms for regulated industries.")
    document.add_heading("Pricing", level=2)
    document.add_paragraph("Pilots start at a fixed fee.", style="List Bullet")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Tier"
    table.rows[0].cells[1].text = "Gold"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class FakeProvider:
    """Stands in for an OpenAI-compatible adapter: vision, summaries, enrichment and embeddings."""

    def __init__(self) -> None:
        self.vision_calls: list[dict] = []
        self.enrichment_batches: list[int] = []
        self.summary_calls = 0
        self.embedded_texts: list[str] = []
        self.fail_embeddings = False

    async def generate_text(self, profile, request):
        images = request.metadata.get("input_images") or []
        if images:
            self.vision_calls.append({"model": profile.models, "images": len(images), "prompt": request.prompt})
            return TextGenerationResult(text="# Scanned page\n\nContract total is 42 EUR, signed by Dana.",
                                        provider="openai_compatible", model="vision-test", input_tokens=900, output_tokens=40)
        if request.response_schema and "notes" in request.response_schema.get("properties", {}):
            ids = [part.split('"')[0] for part in request.prompt.split('<passage id="')[1:]]
            self.enrichment_batches.append(len(ids))
            notes = [{"id": item, "context": f"LLM note for {item}."} for item in ids]
            return TextGenerationResult(text=json.dumps({"notes": notes}), provider="openai_compatible",
                                        model="economy-test", input_tokens=100, output_tokens=20)
        self.summary_calls += 1
        return TextGenerationResult(text="A short test summary of the document.", provider="openai_compatible",
                                    model="economy-test", input_tokens=50, output_tokens=10)

    async def embed(self, profile, request):
        if self.fail_embeddings:
            from app.adapters.base import ProviderExecutionError

            raise ProviderExecutionError("embedding outage")
        self.embedded_texts.extend(request.inputs)
        vectors = []
        for value in request.inputs:
            lowered = value.lower()
            vectors.append([1.0 if "pricing" in lowered or "fee" in lowered else 0.0,
                            1.0 if "contract" in lowered else 0.0, 0.1])
        return EmbeddingResult(vectors=vectors, provider="openai_compatible", model="embed-test", dimensions=3)

    async def test_configuration(self, profile):  # pragma: no cover - not used
        raise NotImplementedError


def configure_providers(client, app, fake: FakeProvider, *, vision: bool = True, text_default: bool = True) -> dict:
    """Create an OpenAI-compatible profile with text + embeddings; optionally set it as vision route."""
    from meetings_contracts import ProviderType

    from app.database import OrganizationAiSettingsRow, LEGACY_ORGANIZATION_ID

    app.state.profile_service.adapters[ProviderType.OPENAI_COMPATIBLE] = fake
    profile = client.post("/v1/provider-profiles", json={
        "name": "Local models", "provider_type": "openai_compatible", "execution_location": "local",
        "base_url": "http://models.local/v1",
        "capabilities": [{"capability": "text_generation", "model": "economy-test"},
                         {"capability": "embeddings", "model": "embed-test"}],
    })
    assert profile.status_code == 201, profile.text
    profile_id = profile.json()["id"]
    assert client.put("/v1/provider-defaults/embeddings", json={
        "policy": "local_only", "local_profile_id": profile_id,
    }).status_code == 200
    if text_default:
        assert client.put("/v1/provider-defaults/text_generation", json={
            "policy": "local_only", "local_profile_id": profile_id,
        }).status_code == 200
    if vision:
        with app.state.database.session_factory.begin() as session:
            session.merge(OrganizationAiSettingsRow(
                organization_id=str(LEGACY_ORGANIZATION_ID), vision_profile_id=profile_id,
                vision_model="vision-test", updated_at=datetime.now(UTC),
            ))
    return profile.json()


def calendar_event(app, user_id, organization_id) -> str:
    from app.database import CalendarEventCacheRow

    event_id = str(uuid4())
    with app.state.database.session_factory.begin() as session:
        session.add(CalendarEventCacheRow(
            id=event_id, organization_id=str(organization_id), user_id=str(user_id), connection_id="conn",
            provider="googlecalendar", event_id=f"evt-{event_id}", starts_at=datetime.now(UTC),
            ends_at=datetime.now(UTC), payload={}, synced_at=datetime.now(UTC),
        ))
    return event_id
